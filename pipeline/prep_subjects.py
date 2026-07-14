#!/usr/bin/env python
"""
Prepare subject list with ages from Excel file for brain extraction.

Reads an Excel file with subject IDs and ages, matches them to available
T1 files, and creates a text file for the SGE array job.

Usage:
    python prep_subjects.py demographics.xlsx --id-col SubjectID --age-col Age
    python prep_subjects.py demographics.xlsx  # tries to auto-detect columns
"""

import os
import sys
import argparse
import glob

def read_excel_or_csv(filepath):
    """Read Excel or CSV file into list of dicts."""
    if filepath.endswith('.csv'):
        import csv
        with open(filepath, 'r') as f:
            reader = csv.DictReader(f)
            return list(reader)
    else:
        # Excel file - try pandas first, fall back to openpyxl
        try:
            import pandas as pd
            df = pd.read_excel(filepath)
            return df.to_dict('records')
        except ImportError:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(filepath)
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                data = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    data.append(dict(zip(headers, row)))
                return data
            except ImportError:
                print("ERROR: Need pandas or openpyxl to read Excel files")
                print("Install with: pip install pandas openpyxl")
                sys.exit(1)


def find_column(data, possible_names):
    """Find a column matching one of the possible names (case-insensitive)."""
    if not data:
        return None
    
    keys = list(data[0].keys())
    for name in possible_names:
        for key in keys:
            if key and name.lower() in key.lower():
                return key
    return None


def main():
    parser = argparse.ArgumentParser(
        description='Prepare subject list with ages from Excel/CSV'
    )
    parser.add_argument('input_file', help='Excel or CSV file with demographics')
    parser.add_argument('--id-col', help='Column name for subject ID')
    parser.add_argument('--age-col', help='Column name for age')
    parser.add_argument('--t1-dir', 
                        default='/ifs/loni/faculty/thompson/four_d/TBI_all/ENIGMA_braininjury/Lesions/T1s',
                        help='Directory with T1 files')
    parser.add_argument('--t1-pattern', default='_T1.nii.gz',
                        help='T1 filename suffix (default: _T1.nii.gz)')
    parser.add_argument('--output', '-o', default='subjects_with_ages.txt',
                        help='Output file (default: subjects_with_ages.txt)')
    
    args = parser.parse_args()
    
    # Read input file
    print(f"Reading: {args.input_file}")
    data = read_excel_or_csv(args.input_file)
    print(f"Found {len(data)} rows")
    
    if not data:
        print("ERROR: No data found in file")
        sys.exit(1)
    
    # Find ID column
    if args.id_col:
        id_col = args.id_col
    else:
        id_col = find_column(data, ['subject', 'subj', 'id', 'participant', 'sub_id', 'subjectid'])
        if not id_col:
            print("ERROR: Could not find subject ID column")
            print(f"Available columns: {list(data[0].keys())}")
            print("Use --id-col to specify")
            sys.exit(1)
    
    print(f"Using ID column: {id_col}")
    
    # Find age column
    if args.age_col:
        age_col = args.age_col
    else:
        age_col = find_column(data, ['age', 'years', 'age_at_scan', 'age_years'])
        if not age_col:
            print("ERROR: Could not find age column")
            print(f"Available columns: {list(data[0].keys())}")
            print("Use --age-col to specify")
            sys.exit(1)
    
    print(f"Using age column: {age_col}")
    
    # Get list of available T1 files
    t1_files = glob.glob(os.path.join(args.t1_dir, f'*{args.t1_pattern}'))
    available_subjects = set()
    for f in t1_files:
        basename = os.path.basename(f)
        sub_id = basename.replace(args.t1_pattern, '')
        available_subjects.add(sub_id)
    
    print(f"Found {len(available_subjects)} T1 files in {args.t1_dir}")
    
    # Match subjects
    matched = []
    missing_t1 = []
    missing_age = []
    
    for row in data:
        sub_id = str(row.get(id_col, '')).strip()
        age = row.get(age_col)
        
        if not sub_id:
            continue
        
        if sub_id not in available_subjects:
            missing_t1.append(sub_id)
            continue
        
        if age is None or (isinstance(age, str) and not age.strip()):
            missing_age.append(sub_id)
            continue
        
        try:
            age_float = float(age)
            matched.append((sub_id, age_float))
        except (ValueError, TypeError):
            print(f"  Warning: Invalid age for {sub_id}: {age}")
            missing_age.append(sub_id)
    
    # Report
    print(f"\nMatched: {len(matched)} subjects")
    
    if missing_t1:
        print(f"Missing T1 files: {len(missing_t1)}")
        if len(missing_t1) <= 10:
            for s in missing_t1:
                print(f"  - {s}")
        else:
            for s in missing_t1[:5]:
                print(f"  - {s}")
            print(f"  ... and {len(missing_t1) - 5} more")
    
    if missing_age:
        print(f"Missing age data: {len(missing_age)}")
        if len(missing_age) <= 10:
            for s in missing_age:
                print(f"  - {s}")
    
    # Age summary
    if matched:
        ages = [a for _, a in matched]
        print(f"\nAge range: {min(ages):.1f} - {max(ages):.1f} years")
        young = sum(1 for a in ages if a <= 16)
        older = sum(1 for a in ages if a > 16)
        print(f"  Age <= 16 (NKI10AndUnder): {young}")
        print(f"  Age > 16 (NKI): {older}")
    
    # Write output
    print(f"\nWriting: {args.output}")
    with open(args.output, 'w') as f:
        for sub_id, age in matched:
            f.write(f"{sub_id}\t{age}\n")
    
    print(f"Done! Created {args.output} with {len(matched)} subjects")
    print(f"\nTo submit jobs:")
    print(f"  qsub -t 1-{len(matched)} run_brain_extraction.sh")


if __name__ == '__main__':
    main()
